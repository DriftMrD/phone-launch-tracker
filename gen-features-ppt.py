#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""9 月国内新机 · 系统应用功能简报（给老板看：做什么、和我们哪条产品有关、来源）。"""
import json
import os
import re
import subprocess
from urllib.parse import urlparse

from lxml import etree
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

REAL_MAP = {}

ROOT = "/Users/shswchengming/WorkBuddy/2026-05-26-10-32-16/deliverables/phone-launch-tracker"
POSTER = "/Users/shswchengming/.cursor/projects/Users-shswchengming-WorkBuddy-2026-05-26-10-32-16-deliverables-phone-launch-tracker/assets/img_v3_0215e_85dd86bd-1d0b-4cd3-a03f-894077f473ag-ea4f47c5-ac2f-4cae-a10d-7b64fa021482.jpg"
ART = "/tmp/ppt-art"
OUT = ROOT + "/2026-09-国内新机-系统应用功能清单.pptx"

W, H = Inches(13.333), Inches(7.5)
BG = RGBColor(0xF6, 0xF3, 0xEE)
PAPER = RGBColor(0xFF, 0xFC, 0xF8)
INK = RGBColor(0x1A, 0x17, 0x14)
MUTED = RGBColor(0x6B, 0x65, 0x60)
FAINT = RGBColor(0x9A, 0x94, 0x8E)
LINE = RGBColor(0xE6, 0xE0, 0xD8)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
ACCENT = RGBColor(0xC4, 0x16, 0x1C)
LINK = RGBColor(0x1A, 0x5F, 0xC2)

DUMP_JS = r"""
const { PHONES, BRAND, PRODUCTS, FEATURES } = require('./app-features-data');
const CAL = [
  { id:'xiaomi18fold', name:'小米 18 Fold', brand:'xiaomi', os:'HyperOS 4', date:'9月7日', status:'已发布', note:'中折叠 · 玄戒 O3' },
  { id:'huaweimatext2', name:'Mate XT 2', brand:'huawei', os:'HarmonyOS 7', date:'9月7日', status:'已发布', note:'三折叠' },
  { id:'puraxview', name:'Pura X View', brand:'huawei', os:'HarmonyOS 7', date:'9月7日', status:'已发布', note:'阔直板' },
  { id:'iphone18', name:'iPhone 18 系列', brand:'apple', os:'iOS 27', date:'9月9日', status:'已发布', note:'Siri AI' },
  { id:'vivox500', name:'vivo X500 系列', brand:'vivo', os:'OriginOS 7', date:'9月16日', status:'已发布', note:'蓝心 4.0' },
  { id:'mate90', name:'Mate 90 系列', brand:'huawei', os:'HarmonyOS 7', date:'9月23日', status:'即将发布', note:'麒麟 9030/9050' },
  { id:'findx10', name:'Find X10 系列', brand:'oppo', os:'ColorOS', date:'9月23日', status:'即将发布', note:'哈苏电影机' },
  { id:'xiaomi18', name:'小米 18 系列', brand:'xiaomi', os:'HyperOS 4', date:'9月24日', status:'即将发布', note:'骁龙 8 Elite Gen6' },
  { id:'magic9', name:'Magic9 系列', brand:'honor', os:'MagicOS 11', date:'9月28日', status:'即将发布', note:'YOYO Agent' }
];
const calIds = new Set(CAL.map(c => c.id));
const calBrands = new Set(CAL.map(c => c.brand));
const osOf = Object.fromEntries(CAL.map(c => [c.brand, c.os]));
const rows = FEATURES.filter(f => {
  if (f.phones && f.phones.length) return f.phones.some(id => calIds.has(id));
  return (f.brands || []).some(b => calBrands.has(b));
}).map(f => {
  const brands = (f.phones && f.phones.length)
    ? [...new Set(PHONES.filter(p => f.phones.includes(p.id)).map(p => p.brand))]
    : (f.brands || []).filter(b => calBrands.has(b));
  const phones = (f.phones && f.phones.length)
    ? PHONES.filter(p => f.phones.includes(p.id) && calIds.has(p.id)).map(p => p.name)
    : CAL.filter(c => brands.includes(c.brand)).map(c => c.name);
  const prod = PRODUCTS.find(p => p.id === f.product);
  const src = f.source || {};
  return {
    name: f.name, product: f.product, productName: prod.name, productColor: prod.color,
    brands, brandLabels: brands.map(b => BRAND[b].label),
    phones: [...new Set(phones)],
    os: [...new Set(brands.map(b => osOf[b]).filter(Boolean))].join(' / '),
    desc: f.desc, note: f.note,
    kind: src.kind || 'update',
    srcLabel: src.label || '',
    url: src.url || ''
  };
});
console.log(JSON.stringify({ CAL, PRODUCTS, BRAND, rows }));
"""

BRAND_PHOTO = {
    "xiaomi": "ph-xiaomi18fold-hero.png",
    "huawei": "ph-matext2-fold.png",
    "apple": "ph-iphone18-hero.png",
    "vivo": "ph-x500-hero.png",
}
BRAND_LOGO = {
    "xiaomi": "xiaomi.png", "huawei": "huawei.png", "apple": "apple.png",
    "vivo": "vivo.png", "oppo": "oppo.png", "honor": "honor.png",
}
BRAND_HEX = {
    "xiaomi": "#ff6900", "huawei": "#c4161c", "apple": "#1d1d1f",
    "vivo": "#415fff", "oppo": "#1a9c2d", "honor": "#1a6fe8",
}


def dump():
    raw = subprocess.check_output(["node", "-e", DUMP_JS], cwd=ROOT)
    return json.loads(raw)


def rasterize():
    os.makedirs(ART, exist_ok=True)
    apps = ROOT + "/images/apps"
    for name in os.listdir(apps):
        if not name.endswith(".svg"):
            continue
        src = os.path.join(apps, name)
        dst = os.path.join(ART, "app-" + name.replace(".svg", ".png"))
        if name == "labs.svg":
            svg = open(src, encoding="utf-8").read().replace("GAME & LABS", "GAME &amp; LABS")
            tmp = "/tmp/labs-fix.svg"
            open(tmp, "w", encoding="utf-8").write(svg)
            src = tmp
        subprocess.check_call(["rsvg-convert", "-w", "1000", src, "-o", dst])
    for name in [
        "iphone18-hero", "iphone18-ios27", "xiaomi18fold-hero", "xiaomi18fold-screen",
        "matext2-fold", "matext2-os", "matext2-pura", "x500-hero", "x500-os",
    ]:
        subprocess.check_call([
            "rsvg-convert", "-w", "1400",
            f"{ROOT}/images/{name}.svg", "-o", f"{ART}/ph-{name}.png",
        ])
    # 无实拍的品牌：Logo 卡
    for bid, hexcol in BRAND_HEX.items():
        rgb = tuple(int(hexcol[i:i + 2], 16) for i in (1, 3, 5))
        im = Image.new("RGB", (1400, 788), rgb)
        logo = Image.open(f"{ROOT}/images/logos/{BRAND_LOGO[bid]}").convert("RGBA")
        logo.thumbnail((380, 380))
        x = (1400 - logo.size[0]) // 2
        y = (788 - logo.size[1]) // 2
        im.paste(logo, (x, y), logo)
        im.save(f"{ART}/logo-{bid}.png")
    im = Image.new("RGB", (1000, 650), (15, 118, 110))
    draw = ImageDraw.Draw(im)
    draw.rounded_rectangle((380, 160, 620, 400), 48, outline=(255, 255, 255), width=8)
    draw.ellipse((455, 215, 545, 305), outline=(255, 255, 255), width=8)
    im.save(f"{ART}/app-passwords.png")


def rgb_hex(h):
    h = h.lstrip("#")
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def set_ea(run, name):
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rPr, qn(tag))
        el.set("typeface", name)


def font(run, name, size, bold=False, color=INK):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    set_ea(run, name)


def box(slide, l, t, w, h, fill=None, line=None):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, h)
    sh.shadow.inherit = False
    sh.line.fill.background()
    if line:
        sh.line.fill.solid()
        sh.line.color.rgb = line
        sh.line.width = Pt(0.75)
    if fill:
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
    else:
        sh.fill.background()
    return sh


def roundbox(slide, l, t, w, h, fill=PAPER, line=LINE):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
    sh.shadow.inherit = False
    sh.adjustments[0] = 0.06
    if fill:
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
    else:
        sh.fill.background()
    if line:
        sh.line.fill.solid()
        sh.line.color.rgb = line
        sh.line.width = Pt(0.75)
    else:
        sh.line.fill.background()
    return sh


def tx(slide, l, t, w, h, text, name="PingFang SC", size=14, bold=False,
       color=INK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spc=1.12):
    sh = slide.shapes.add_textbox(l, t, w, h)
    tf = sh.text_frame
    tf.word_wrap = True
    try:
        tf.vertical_anchor = anchor
    except Exception:
        pass
    p = tf.paragraphs[0]
    p.alignment = align
    p.space_before = Pt(0)
    p.space_after = Pt(0)
    p.line_spacing = spc
    run = p.add_run()
    run.text = text
    font(run, name, size, bold, color)
    return sh


def footer(slide, page, total, dark=False):
    c = RGBColor(0xC8, 0xC0, 0xB6) if not dark else RGBColor(0x8A, 0x82, 0x7A)
    tx(slide, Inches(0.5), Inches(7.18), Inches(9.2), Inches(0.24),
       "2026年9月竞品系统应用  ·  功能简报", size=10, color=c)
    tx(slide, Inches(11.3), Inches(7.18), Inches(1.5), Inches(0.24),
       f"{page}  /  {total}", size=10, color=c, align=PP_ALIGN.RIGHT)


def pic(slide, path, l, t, w, h):
    if path and os.path.exists(path):
        slide.shapes.add_picture(path, l, t, w, h)


def domain_of(url):
    try:
        host = urlparse(url).hostname or url
        return host.replace("www.", "")
    except Exception:
        return url


def link_shape(shape, url):
    if url:
        shape.click_action.hyperlink.address = url


def new(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    return s


def bg(s):
    box(s, 0, 0, W, H, BG)
    box(s, 0, 0, Inches(0.12), H, ACCENT)


def cover(prs, data, total):
    s = new(prs)
    box(s, 0, 0, W, H, RGBColor(0x16, 0x14, 0x12))
    box(s, 0, 0, Inches(0.16), H, ACCENT)
    tx(s, Inches(0.7), Inches(1.45), Inches(12), Inches(0.32),
       "2026 年 9 月国内新机发布季", size=13, bold=True, color=RGBColor(0xE8, 0xB4, 0x3A))
    tx(s, Inches(0.65), Inches(1.9), Inches(12), Inches(1.4),
       "竞品系统应用新功能\n简报", name="Songti SC", size=44, bold=True, color=WHITE, spc=1.02)
    tx(s, Inches(0.7), Inches(3.7), Inches(11), Inches(0.7),
       "最近各家发布了哪些系统应用能力、和我们哪条产品有关、出处在哪。一页讲清一个功能。",
       size=16, color=RGBColor(0xC8, 0xC0, 0xB6), spc=1.3)
    items = [
        ("9 台新机", "小米 / 华为 / 苹果 / vivo / OPPO / 荣耀"),
        (f"{len(data['rows'])} 个功能", "对照我们 15 个系统应用"),
        ("每条带来源", "官方或媒体原文可点开"),
    ]
    x = Inches(0.7)
    for a, b in items:
        roundbox(s, x, Inches(5.05), Inches(3.7), Inches(1.15),
                 fill=RGBColor(0x24, 0x21, 0x1E), line=RGBColor(0x3A, 0x34, 0x30))
        tx(s, x + Inches(0.22), Inches(5.18), Inches(3.3), Inches(0.42), a, size=18, bold=True, color=WHITE)
        tx(s, x + Inches(0.22), Inches(5.62), Inches(3.3), Inches(0.4), b, size=12, color=RGBColor(0xB8, 0xB0, 0xA6))
        x += Inches(3.9)
    footer(s, 1, total, dark=True)


def agenda(prs, data, total):
    s = new(prs)
    bg(s)
    tx(s, Inches(0.55), Inches(0.32), Inches(12), Inches(0.26), "简报怎么用", size=11, bold=True, color=ACCENT)
    tx(s, Inches(0.5), Inches(0.55), Inches(12), Inches(0.5),
       "三件事", name="Songti SC", size=28, bold=True)
    cards = [
        ("01", "发了什么", "9 月国内旗舰各自带上哪套系统，系统应用上新增了哪些可感知的能力。"),
        ("02", "和我们哪条产品有关", "按录音机、文件管理、记事本、日历、相册、语音助手等我们自己的产品线对齐，不讲芯片和影像硬件。"),
        ("03", "是不是真的", "每条功能写明出现在哪台机、哪套系统，并附上官方或媒体原文链接，点开可核。"),
    ]
    y = Inches(1.4)
    for n, t, d in cards:
        roundbox(s, Inches(0.5), y, Inches(12.3), Inches(1.55), PAPER, LINE)
        tx(s, Inches(0.8), y + Inches(0.35), Inches(1.1), Inches(0.7), n, name="Songti SC", size=28, bold=True, color=ACCENT)
        tx(s, Inches(2.1), y + Inches(0.28), Inches(10.2), Inches(0.4), t, size=20, bold=True)
        tx(s, Inches(2.1), y + Inches(0.78), Inches(10.2), Inches(0.55), d, size=15, color=MUTED, spc=1.25)
        y += Inches(1.7)
    footer(s, 2, total)


def phones(prs, data, total):
    s = new(prs)
    bg(s)
    tx(s, Inches(0.55), Inches(0.28), Inches(8), Inches(0.24), "本月新机", size=11, bold=True, color=ACCENT)
    tx(s, Inches(0.5), Inches(0.5), Inches(8.5), Inches(0.48),
       "9 月国内发布会机型", name="Songti SC", size=26, bold=True)
    if os.path.exists(POSTER):
        pic(s, POSTER, Inches(9.55), Inches(1.35), Inches(3.3), Inches(5.4))
    cols = [("日期", 1.2), ("品牌", 0.95), ("机型", 2.2), ("系统", 1.7), ("状态", 1.15)]
    x0, y = Inches(0.5), Inches(1.2)
    box(s, x0, y, Inches(8.7), Inches(0.38), RGBColor(0x16, 0x14, 0x12))
    x = x0 + Inches(0.12)
    for name, w in cols:
        tx(s, x, y, Inches(w), Inches(0.38), name, size=12, bold=True, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
        x += Inches(w)
    y = Inches(1.62)
    for i, c in enumerate(data["CAL"]):
        fill = PAPER if i % 2 == 0 else RGBColor(0xF3, 0xEE, 0xE6)
        box(s, x0, y, Inches(8.7), Inches(0.5), fill)
        box(s, x0, y, Inches(0.08), Inches(0.5), rgb_hex(data["BRAND"][c["brand"]]["color"]))
        vals = [c["date"], data["BRAND"][c["brand"]]["label"], c["name"], c["os"], c["status"]]
        x = x0 + Inches(0.12)
        for (name, w), v in zip(cols, vals):
            col = ACCENT if v == "即将发布" else INK
            tx(s, x, y, Inches(w), Inches(0.5), v, size=13, bold=(name == "机型"), color=col, anchor=MSO_ANCHOR.MIDDLE)
            x += Inches(w)
        y += Inches(0.5)
    tx(s, Inches(0.55), Inches(6.3), Inches(8.6), Inches(0.6),
       "9 月 15 日还有一场联发科天玑 9600 芯片发布，首发落在 vivo X500 与 Find X10 上，功能记在这两台手机里。",
       size=13, color=MUTED, spc=1.2)
    footer(s, 3, total)


def our_products(prs, data, total):
    s = new(prs)
    bg(s)
    tx(s, Inches(0.55), Inches(0.28), Inches(12), Inches(0.24), "我们对照的产品", size=11, bold=True, color=ACCENT)
    tx(s, Inches(0.5), Inches(0.5), Inches(12), Inches(0.48),
       "15 个系统应用", name="Songti SC", size=26, bold=True)
    tx(s, Inches(0.55), Inches(1.02), Inches(12), Inches(0.32),
       "后面每一章对应我们自己的一条产品。竞品做了什么，直接看对我们意味着什么。", size=14, color=MUTED)
    counts = {p["id"]: 0 for p in data["PRODUCTS"]}
    for r in data["rows"]:
        counts[r["product"]] += 1
    for i, p in enumerate(data["PRODUCTS"]):
        col, row = i % 5, i // 5
        x = Inches(0.5) + col * Inches(2.55)
        y = Inches(1.5) + row * Inches(1.75)
        roundbox(s, x, y, Inches(2.42), Inches(1.58), PAPER, LINE)
        box(s, x, y, Inches(0.1), Inches(1.58), rgb_hex(p["color"]))
        icon = f"{ART}/app-{p['id']}.png"
        if os.path.exists(icon):
            pic(s, icon, x + Inches(0.22), y + Inches(0.14), Inches(0.7), Inches(0.45))
        tx(s, x + Inches(0.22), y + Inches(0.68), Inches(2.05), Inches(0.36),
           p["name"], size=16, bold=True)
        n = counts[p["id"]]
        tx(s, x + Inches(0.22), y + Inches(1.08), Inches(2.05), Inches(0.32),
           f"{n} 个功能" if n else "本月未见新能力", size=12, color=MUTED)
    footer(s, 4, total)


def chapter(prs, prod, feats, page, total):
    s = new(prs)
    c = rgb_hex(prod["color"])
    box(s, 0, 0, W, H, RGBColor(0x16, 0x14, 0x12))
    box(s, 0, 0, Inches(0.16), H, c)
    tx(s, Inches(0.7), Inches(1.7), Inches(12), Inches(0.3),
       "系统应用", size=13, bold=True, color=c)
    tx(s, Inches(0.65), Inches(2.1), Inches(12), Inches(0.7),
       prod["name"], name="Songti SC", size=48, bold=True, color=WHITE)
    tx(s, Inches(0.7), Inches(2.95), Inches(12), Inches(0.4),
       f"本月竞品 {len(feats)} 个相关功能，下一页起一页一个。", size=16, color=RGBColor(0xC8, 0xC0, 0xB6))
    names = "  ·  ".join(f["name"] for f in feats)
    tx(s, Inches(0.7), Inches(3.6), Inches(12), Inches(2.8), names, size=15, color=RGBColor(0xB8, 0xB0, 0xA6), spc=1.45)
    footer(s, page, total, dark=True)


def feature_slide(prs, data, feat, idx, n, page, total):
    s = new(prs)
    bg(s)
    pc = rgb_hex(feat["productColor"])
    box(s, 0, 0, Inches(0.12), H, pc)

    tx(s, Inches(0.5), Inches(0.22), Inches(12.3), Inches(0.26),
       f"{feat['productName']}    {idx} / {n}", size=12, bold=True, color=pc)
    tx(s, Inches(0.48), Inches(0.46), Inches(12.3), Inches(0.62),
       feat["name"], name="Songti SC", size=26, bold=True, color=INK)

    # 左栏：出处真实图或来源页截图
    info = REAL_MAP.get(feat["name"]) or {}
    art = info.get("path")
    left, top, lw, lh = Inches(0.42), Inches(1.16), Inches(4.6), Inches(5.78)
    roundbox(s, left, top, lw, lh, PAPER, LINE)
    if art and os.path.exists(art):
        pic(s, art, left + Inches(0.08), top + Inches(0.08), lw - Inches(0.16), lh - Inches(0.42))
    cap = "出处配图" if info.get("via") == "source-ui" else "来源页截图"
    tx(s, left + Inches(0.16), top + lh - Inches(0.38), lw - Inches(0.32), Inches(0.28),
       cap, size=11, color=MUTED)

    # 右栏
    rx, rw = Inches(5.2), Inches(7.55)
    y = Inches(1.18)
    tx(s, rx, y, rw, Inches(0.28), "这是什么", size=12, bold=True, color=pc)
    tx(s, rx, y + Inches(0.3), rw, Inches(1.7), feat["desc"], size=15, color=INK, spc=1.28)

    y = Inches(3.22)
    tx(s, rx, y, rw, Inches(0.26), "出现在", size=12, bold=True, color=pc)
    phones = "  ·  ".join(feat["phones"]) if feat["phones"] else " / ".join(feat["brandLabels"])
    tx(s, rx, y + Inches(0.28), rw, Inches(0.55),
       f"{' / '.join(feat['brandLabels'])}    {feat['os']}\n{phones}",
       size=14, color=INK, spc=1.2)

    y = Inches(4.15)
    tx(s, rx, y, rw, Inches(0.26), "和我们产品的关系", size=12, bold=True, color=pc)
    note = re.sub(r"tOS\s*", "我们", feat.get("note") or "")
    tx(s, rx, y + Inches(0.28), rw, Inches(1.55), note, size=14, color=INK, spc=1.25)

    y = Inches(6.05)
    card = roundbox(s, rx, y, rw, Inches(0.85), PAPER, LINE)
    tx(s, rx + Inches(0.18), y + Inches(0.08), Inches(1.1), Inches(0.28), "来源", size=11, bold=True, color=pc)
    src_title = feat["srcLabel"] or "发布会 / 系统更新"
    url = feat.get("url") or ""
    line = (domain_of(url) + "  ·  " + src_title) if url else src_title
    sh = s.shapes.add_textbox(rx + Inches(0.18), y + Inches(0.34), Inches(5.6), Inches(0.42))
    tf = sh.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.space_before = Pt(0)
    p.space_after = Pt(0)
    run = p.add_run()
    run.text = line
    font(run, "PingFang SC", 13, True, LINK if url else INK)
    if url:
        run.font.underline = True
        run.hyperlink.address = url
        btn = roundbox(s, rx + Inches(5.85), y + Inches(0.22), Inches(1.5), Inches(0.42),
                       fill=RGBColor(0x1A, 0x5F, 0xC2), line=RGBColor(0x1A, 0x5F, 0xC2))
        tx(s, rx + Inches(5.85), y + Inches(0.26), Inches(1.5), Inches(0.36),
           "打开原文", size=12, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        link_shape(card, url)
        link_shape(btn, url)
        hit = box(s, rx, y, rw, Inches(0.85), fill=None, line=None)
        hit.fill.background()
        link_shape(hit, url)

    footer(s, page, total)


def close_slide(prs, data, page, total):
    s = new(prs)
    bg(s)
    tx(s, Inches(0.55), Inches(0.28), Inches(12), Inches(0.24), "收口", size=11, bold=True, color=ACCENT)
    tx(s, Inches(0.5), Inches(0.5), Inches(12), Inches(0.5),
       "建议优先看这几条", name="Songti SC", size=26, bold=True)
    items = [
        ("录音机", "录完就能问、大纲可跳转、摘要不要跳出 App。小米会议大纲、OPPO 模板、荣耀端侧纪要都在打同一条路。"),
        ("文件管理", "回收站是底线；下一步是「用一句话找到文件」和「256GB 机能多出 20GB」。"),
        ("记事本 / 日历", "看见即记下、截图进日程、录音抽出待办。入口越短越像系统能力。"),
        ("语音助手", "跨应用办事 + 长期记忆。小米开始走订阅，荣耀给系统级权限，苹果大陆暂不上 Siri AI。"),
        ("扫一扫 / 换机", "锁屏或外屏直接扫码；换机要搬走闹钟、笔记、录音和桌面，而不只是照片。"),
    ]
    y = Inches(1.2)
    for t, d in items:
        roundbox(s, Inches(0.5), y, Inches(12.3), Inches(1.05), PAPER, LINE)
        tx(s, Inches(0.75), y + Inches(0.14), Inches(11.8), Inches(0.3), t, size=16, bold=True, color=ACCENT)
        tx(s, Inches(0.75), y + Inches(0.48), Inches(11.8), Inches(0.48), d, size=14, color=INK)
        y += Inches(1.12)
    footer(s, page, total)


def main():
    global REAL_MAP
    rasterize()
    mp = "/tmp/ppt-real/final/map.json"
    if os.path.exists(mp):
        REAL_MAP = json.load(open(mp, encoding="utf-8"))
    data = dump()
    by = {p["id"]: [] for p in data["PRODUCTS"]}
    for r in data["rows"]:
        by[r["product"]].append(r)
    chapters = [(p, by[p["id"]]) for p in data["PRODUCTS"] if by[p["id"]]]

    n_feat = sum(len(fs) for _, fs in chapters)
    total = 4 + len(chapters) + n_feat + 1

    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    cover(prs, data, total)
    agenda(prs, data, total)
    phones(prs, data, total)
    our_products(prs, data, total)

    page = 5
    for prod, feats in chapters:
        chapter(prs, prod, feats, page, total)
        page += 1
        for i, f in enumerate(feats, 1):
            feature_slide(prs, data, f, i, len(feats), page, total)
            page += 1
    close_slide(prs, data, page, total)

    prs.save(OUT)
    print("wrote", OUT)
    print("slides", len(prs.slides), "features", n_feat)


if __name__ == "__main__":
    main()
